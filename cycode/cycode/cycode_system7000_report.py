import json
import os
import re
import shutil
import site
import subprocess
import sys
import time
from collections import Counter, defaultdict
from datetime import datetime
from pathlib import Path

PROJECT_ID = 36416
STATUS = "Open"
PAGE_SIZE = 1000
OUTPUT_FILE = "OneLink_Cycode_Report.xlsx"
SEVERITIES = ("High", "Critical", "Medium", "Low", "Info")

# Cycode can return historical/legacy repository names in violation records even
# when the GUI shows only the current repository name. Normalize those aliases
# so the Excel matches the GUI repository selector and totals.
REPOSITORY_ALIASES = {
    "s7k-build-pipelines": "cxs-s7k-build-pipelines",
}

ANSI_RE = re.compile(r"\x1B(?:[@-Z\\-_]|\[[0-?]*[ -/]*[@-~])")


def find_cycode_executable():
    """Find cycode.exe from PATH or the current Python user's Scripts folder."""
    found = shutil.which("cycode")
    if found:
        return found

    user_site = Path(site.getusersitepackages())
    candidate = user_site.parent / "Scripts" / "cycode.exe"
    if candidate.exists():
        return str(candidate)

    appdata = os.environ.get("APPDATA")
    if appdata:
        py_tag = f"Python{sys.version_info.major}{sys.version_info.minor}"
        candidate = Path(appdata) / "Python" / py_tag / "Scripts" / "cycode.exe"
        if candidate.exists():
            return str(candidate)

    raise FileNotFoundError(
        "Cycode CLI was not found. Install it with:\n"
        "  py -m pip install --upgrade cycode\n"
        "Then authenticate with:\n"
        "  cycode auth"
    )


def extract_json(text):
    """Extract the Cycode JSON object even if Rich/logging text appears around it."""
    cleaned = ANSI_RE.sub("", text or "")
    decoder = json.JSONDecoder()

    for match in re.finditer(r"[\{\[]", cleaned):
        try:
            obj, _ = decoder.raw_decode(cleaned[match.start():])
        except json.JSONDecodeError:
            continue

        if isinstance(obj, dict) and (
            "items" in obj or "next_page_token" in obj or "page_size" in obj
        ):
            return obj

    raise ValueError("Could not find a valid Cycode JSON response in the CLI output.")


def run_cycode_json(cycode_exe, next_page_token=None, retries=3):
    cmd = [
        cycode_exe,
        "platform",
        "violations",
        "list",
        "--status",
        STATUS,
        "--project-ids",
        str(PROJECT_ID),
        "--page-size",
        str(PAGE_SIZE),
    ]

    if next_page_token:
        cmd += ["--next-page-token", next_page_token]

    last_error = None

    for attempt in range(1, retries + 1):
        result = subprocess.run(
            cmd,
            capture_output=True,
            text=True,
            encoding="utf-8",
            errors="replace",
        )

        if result.returncode == 0:
            try:
                return extract_json(result.stdout)
            except Exception as exc:
                last_error = exc
        else:
            combined = f"{result.stdout}\n{result.stderr}"
            if "credentials" in combined.lower() or "cycode auth" in combined.lower():
                raise RuntimeError(
                    "Cycode authentication is required. Run 'cycode auth' and then run this script again."
                )
            last_error = RuntimeError(combined.strip())

        if attempt < retries:
            print(f"Page request failed (attempt {attempt}/{retries}). Retrying...")
            time.sleep(attempt * 2)

    raise RuntimeError(f"Cycode request failed after {retries} attempts: {last_error}")


def find_first_key(obj, wanted_key):
    """Fallback search for a key anywhere in a nested JSON object."""
    if isinstance(obj, dict):
        if wanted_key in obj and obj[wanted_key] not in (None, ""):
            return obj[wanted_key]
        for value in obj.values():
            found = find_first_key(value, wanted_key)
            if found not in (None, ""):
                return found
    elif isinstance(obj, list):
        for value in obj:
            found = find_first_key(value, wanted_key)
            if found not in (None, ""):
                return found
    return None


def normalize_repository_name(repo):
    if not repo:
        return "(No repository)"
    value = str(repo).strip()
    return REPOSITORY_ALIASES.get(value.casefold(), value)


def get_repo_and_branch(item):
    details = item.get("detection_details")
    repo = None
    branch = None

    if isinstance(details, dict):
        repo = details.get("repository_name")
        branch = details.get("branch_name")
    elif isinstance(details, list):
        for detail in details:
            if not isinstance(detail, dict):
                continue
            repo = repo or detail.get("repository_name")
            branch = branch or detail.get("branch_name")
            if repo and branch:
                break

    # Fallback in case Cycode changes the nesting slightly.
    repo = repo or find_first_key(item, "repository_name")
    branch = branch or find_first_key(item, "branch_name")

    repo = normalize_repository_name(repo)
    branch = str(branch).strip() if branch else "(No branch)"
    return repo, branch


def normalize_severity(item):
    value = item.get("severity") or item.get("risk_score_severity")
    if value is None:
        value = find_first_key(item, "severity")
    if value is None:
        raise ValueError(f"Severity is missing for violation {item.get('id', '<unknown>')}")

    normalized = str(value).strip().title()
    allowed = {s.lower(): s for s in SEVERITIES}
    if normalized.lower() not in allowed:
        raise ValueError(
            f"Unexpected severity '{value}' for violation {item.get('id', '<unknown>')}"
        )
    return allowed[normalized.lower()]


def collect_counts(cycode_exe):
    repo_counts = defaultdict(Counter)
    branch_counts = defaultdict(Counter)
    seen_ids = set()

    token = None
    page_no = 0
    total_items = 0
    duplicate_items = 0

    while True:
        page_no += 1
        data = run_cycode_json(cycode_exe, token)
        items = data.get("items") or []

        if not isinstance(items, list):
            raise ValueError("Cycode response field 'items' is not a list.")

        added_this_page = 0
        for item in items:
            if not isinstance(item, dict):
                continue

            violation_id = item.get("id")
            if violation_id:
                if violation_id in seen_ids:
                    duplicate_items += 1
                    continue
                seen_ids.add(violation_id)

            repo, branch = get_repo_and_branch(item)
            severity = normalize_severity(item)

            repo_counts[repo][severity] += 1
            branch_counts[(repo, branch)][severity] += 1
            total_items += 1
            added_this_page += 1

        print(
            f"Page {page_no}: received {len(items):,}, "
            f"added {added_this_page:,}, total {total_items:,}"
        )

        token = data.get("next_page_token")
        if not token:
            break

    if duplicate_items:
        print(f"Skipped {duplicate_items:,} duplicate violation IDs.")

    return repo_counts, branch_counts, total_items


def require_openpyxl():
    try:
        import openpyxl  # noqa: F401
    except ModuleNotFoundError:
        raise SystemExit(
            "openpyxl is required to create Excel. Install it once with:\n"
            "  py -m pip install openpyxl"
        )


def write_workbook(repo_counts, branch_counts, total_items, output_path):
    require_openpyxl()

    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Repository Summary"
    bs = wb.create_sheet("Branch Summary")

    try:
        wb.calculation.fullCalcOnLoad = True
        wb.calculation.forceFullCalc = True
        wb.calculation.calcMode = "auto"
    except Exception:
        pass

    title_fill = PatternFill("solid", fgColor="17365D")
    header_fill = PatternFill("solid", fgColor="D9EAF7")
    total_fill = PatternFill("solid", fgColor="E2F0D9")
    formula_fill = PatternFill("solid", fgColor="F2F2F2")
    white_font = Font(color="FFFFFF", bold=True, size=14)
    header_font = Font(bold=True, color="1F1F1F")
    imported_font = Font(color="008000")
    formula_font = Font(color="000000")
    static_font = Font(color="666666")
    thin_blue = Side(style="thin", color="5B9BD5")
    total_border = Border(top=Side(style="thin", color="7F7F7F"))

    headers = [
        "Repository",
        "High",
        "Critical",
        "Medium",
        "Low",
        "Info",
        "C+H+M+L",
        "C+H+M+L+Info",
    ]

    def setup_sheet(sheet, title, subtitle, headers_list):
        sheet.sheet_view.showGridLines = False
        sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=len(headers_list))
        cell = sheet.cell(1, 1, title)
        cell.fill = title_fill
        cell.font = white_font
        cell.alignment = Alignment(horizontal="left", vertical="center")
        sheet.row_dimensions[1].height = 26

        sheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=len(headers_list))
        meta = sheet.cell(2, 1, subtitle)
        meta.font = static_font

        for col, header in enumerate(headers_list, 1):
            c = sheet.cell(4, col, header)
            c.fill = header_fill
            c.font = header_font
            c.alignment = Alignment(horizontal="center", vertical="center")
            c.border = Border(bottom=thin_blue)

        sheet.freeze_panes = "A5"

    generated = datetime.now().strftime("%Y-%m-%d %H:%M:%S")
    subtitle = f"Project ID: {PROJECT_ID} | Status: {STATUS} | Generated: {generated}"

    setup_sheet(ws, "System 7000 - Cycode Open Violations", subtitle, headers)

    sorted_repos = sorted(
        repo_counts.items(),
        key=lambda x: (-sum(x[1].values()), x[0].lower()),
    )

    start_row = 5
    for idx, (repo, counts) in enumerate(sorted_repos, start_row):
        ws.cell(idx, 1, repo)
        ws.cell(idx, 2, counts.get("High", 0))
        ws.cell(idx, 3, counts.get("Critical", 0))
        ws.cell(idx, 4, counts.get("Medium", 0))
        ws.cell(idx, 5, counts.get("Low", 0))
        ws.cell(idx, 6, counts.get("Info", 0))
        ws.cell(idx, 7, f"=SUM(B{idx}:E{idx})")
        ws.cell(idx, 8, f"=SUM(B{idx}:F{idx})")

        for col in range(1, 7):
            ws.cell(idx, col).font = imported_font
        for col in range(7, 9):
            ws.cell(idx, col).font = formula_font
            ws.cell(idx, col).fill = formula_fill

    last_data_row = start_row + len(sorted_repos) - 1
    total_row = last_data_row + 1
    ws.cell(total_row, 1, "TOTAL")
    ws.cell(total_row, 1).font = Font(bold=True)
    for col in range(2, 9):
        letter = get_column_letter(col)
        ws.cell(total_row, col, f"=SUM({letter}{start_row}:{letter}{last_data_row})")
        ws.cell(total_row, col).font = Font(bold=True)
    for col in range(1, 9):
        ws.cell(total_row, col).fill = total_fill
        ws.cell(total_row, col).border = total_border

    ws.auto_filter.ref = f"A4:H{last_data_row}"
    ws.column_dimensions["A"].width = 42
    for col in "BCDEFGH":
        ws.column_dimensions[col].width = 16
    for row in ws.iter_rows(min_row=5, max_row=total_row, min_col=2, max_col=8):
        for cell in row:
            cell.number_format = "#,##0"
            cell.alignment = Alignment(horizontal="right")

    branch_headers = [
        "Repository",
        "Branch",
        "High",
        "Critical",
        "Medium",
        "Low",
        "Info",
        "C+H+M+L",
        "C+H+M+L+Info",
    ]
    setup_sheet(bs, "System 7000 - Branch-wise Open Violations", subtitle, branch_headers)

    # Match the Cycode GUI branch dropdown for the project-wide view:
    # - keep every real/named branch returned by Cycode, regardless of severity
    # - exclude only records where Cycode has no branch name
    # This is important because the GUI branch selector can include Info-only branches too.
    gui_branch_counts = {
        key: counts
        for key, counts in branch_counts.items()
        if key[1] != "(No branch)"
    }

    unique_branch_names = {branch for (_, branch) in gui_branch_counts}
    named_branch_violations = sum(sum(counts.values()) for counts in gui_branch_counts.values())
    no_branch_violations = sum(
        sum(counts.values())
        for (repo, branch), counts in branch_counts.items()
        if branch == "(No branch)"
    )

    sorted_branches = sorted(
        gui_branch_counts.items(),
        key=lambda x: (x[0][0].lower(), -sum(x[1].values()), x[0][1].lower()),
    )

    for idx, ((repo, branch), counts) in enumerate(sorted_branches, start_row):
        bs.cell(idx, 1, repo)
        bs.cell(idx, 2, branch)
        bs.cell(idx, 3, counts.get("High", 0))
        bs.cell(idx, 4, counts.get("Critical", 0))
        bs.cell(idx, 5, counts.get("Medium", 0))
        bs.cell(idx, 6, counts.get("Low", 0))
        bs.cell(idx, 7, counts.get("Info", 0))
        bs.cell(idx, 8, f"=SUM(C{idx}:F{idx})")
        bs.cell(idx, 9, f"=SUM(C{idx}:G{idx})")

        for col in range(1, 8):
            bs.cell(idx, col).font = imported_font
        for col in range(8, 10):
            bs.cell(idx, col).font = formula_font
            bs.cell(idx, col).fill = formula_fill

    branch_last_row = start_row + len(sorted_branches) - 1
    branch_total_row = branch_last_row + 1
    bs.cell(branch_total_row, 1, "TOTAL")
    bs.cell(branch_total_row, 1).font = Font(bold=True)
    for col in range(3, 10):
        letter = get_column_letter(col)
        bs.cell(branch_total_row, col, f"=SUM({letter}{start_row}:{letter}{branch_last_row})")
        bs.cell(branch_total_row, col).font = Font(bold=True)
    for col in range(1, 10):
        bs.cell(branch_total_row, col).fill = total_fill
        bs.cell(branch_total_row, col).border = total_border

    bs.auto_filter.ref = f"A4:I{branch_last_row}"
    bs.column_dimensions["A"].width = 42
    bs.column_dimensions["B"].width = 45
    for col in "CDEFGHI":
        bs.column_dimensions[col].width = 16
    for row in bs.iter_rows(min_row=5, max_row=branch_total_row, min_col=3, max_col=9):
        for cell in row:
            cell.number_format = "#,##0"
            cell.alignment = Alignment(horizontal="right")

    # Store API-derived checks in the visible notes.
    ws.cell(2, 1).value = subtitle + f" | Violations fetched: {total_items:,}"
    bs.cell(2, 1).value = (
        subtitle
        + f" | Named branch names: {len(unique_branch_names):,}"
        + f" | Repo/branch rows: {len(gui_branch_counts):,}"
        + f" | Named-branch violations: {named_branch_violations:,}"
        + f" | No-branch violations excluded: {no_branch_violations:,}"
    )

    wb.save(output_path)


def main():
    print("System 7000 Cycode report")
    print(f"Project ID: {PROJECT_ID} | Status: {STATUS}")

    cycode_exe = find_cycode_executable()
    print(f"Cycode CLI: {cycode_exe}")

    repo_counts, branch_counts, total_items = collect_counts(cycode_exe)

    repo_total = sum(sum(c.values()) for c in repo_counts.values())
    branch_total = sum(sum(c.values()) for c in branch_counts.values())

    if repo_total != total_items or branch_total != total_items:
        raise RuntimeError(
            f"Validation failed: fetched={total_items}, repo_total={repo_total}, branch_total={branch_total}"
        )

    output_path = Path.cwd() / OUTPUT_FILE
    write_workbook(repo_counts, branch_counts, total_items, output_path)

    print("\nCompleted successfully.")
    named_pairs = {key: counts for key, counts in branch_counts.items() if key[1] != "(No branch)"}
    unique_named_branches = {branch for (_, branch) in named_pairs}
    named_branch_violations = sum(sum(c.values()) for c in named_pairs.values())
    no_branch_violations = total_items - named_branch_violations

    print(f"Repositories: {len(repo_counts):,}")
    print(f"Named repository/branch combinations: {len(named_pairs):,}")
    print(f"Unique named branches: {len(unique_named_branches):,}")
    print(f"Violations on named branches: {named_branch_violations:,}")
    print(f"Violations with no branch: {no_branch_violations:,}")
    print(f"Open violations: {total_items:,}")
    print(f"Excel: {output_path}")


if __name__ == "__main__":
    try:
        main()
    except KeyboardInterrupt:
        print("\nStopped by user.")
        sys.exit(130)
    except Exception as exc:
        print(f"\nERROR: {exc}")
        sys.exit(1)

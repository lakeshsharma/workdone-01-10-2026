"""
Build a per-studio Excel workbook mirroring the template
(nCompass_MQ_FY26_JQ-26.xlsx): Master Sheet, ActiveBranches, VulnDetails.

VulnDetails has two side-by-side blocks — prior quarter (e.g. MQ'26) filled from
saved history, and current quarter (e.g. JQ'26) from this run.

`rows` is a list of dicts:
    {"project": str, "key": str, "branch": str, "c": int, "h": int, "m": int, "l": int}
`history` is a scope.History used to look up the prior-quarter numbers.
"""
from __future__ import annotations

import os
from collections import OrderedDict
from datetime import datetime

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from . import config

_HEADER_FILL = PatternFill("solid", fgColor=config.COLOR_HEADER)
_TOTAL_FILL = PatternFill("solid", fgColor=config.COLOR_TOTAL)
_YELLOW = PatternFill("solid", fgColor=config.COLOR_YELLOW)
_CRIT = PatternFill("solid", fgColor=config.COLOR_CRIT)
_BOLD = Font(bold=True)
_CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
_thin = Side(style="thin", color="BFBFBF")
_BORDER = Border(left=_thin, right=_thin, top=_thin, bottom=_thin)


def _sheet_name(label: str, suffix: str) -> str:
    """Excel sheet names: max 31 chars, no  : \\ / ? * [ ] ."""
    name = f"{label}-{suffix}"
    for ch in ':\\/?*[]':
        name = name.replace(ch, " ")
    if len(name) > 31:                       # keep the suffix readable
        keep = 31 - len(suffix) - 1
        name = f"{label[:max(keep, 1)]}-{suffix}"[:31]
    return name


def _style_header(ws, row, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row, c)
        cell.fill = _HEADER_FILL
        cell.font = _BOLD
        cell.alignment = _CENTER
        cell.border = _BORDER


def _autofit(ws, widths):
    for col, w in widths.items():
        ws.column_dimensions[col].width = w


def _projects_in_order(rows):
    """OrderedDict project -> {key, branches:[...]} preserving first-seen order."""
    out: "OrderedDict[str, dict]" = OrderedDict()
    for r in rows:
        p = out.setdefault(r["project"], {"key": r["key"], "branches": []})
        if r["branch"] not in p["branches"]:
            p["branches"].append(r["branch"])
    return out


def _master_sheet(wb, quarter, projects):
    ws = wb.create_sheet(_sheet_name(quarter, "Master Sheet"))
    headers = ["No.", "Product", "Applicable for Data", "Repo Link",
               "Sonar Report Available?", "SonarQube Project",
               "SonarQube Project Link", "Comments"]
    ws.append(headers)
    _style_header(ws, 1, len(headers))
    for i, (name, info) in enumerate(projects.items(), 1):
        link = f"https://{config.SONAR_HOST}/dashboard?id={info['key']}"
        ws.append([i, name, "Yes", "", "Yes", info["key"], link, ""])
    _autofit(ws, {"A": 6, "B": 34, "C": 14, "D": 18, "E": 12, "F": 34, "G": 55, "H": 20})
    ws.freeze_panes = "A2"


def _active_branches_sheet(wb, quarter, projects):
    ws = wb.create_sheet(_sheet_name(quarter, "ActiveBranches"))
    headers = ["No.", "Product", "Applicable for Data", "Impacted Product",
               "Active Branch1", "Active Branch2", "Active Branch3", "Comments"]
    ws.append(headers)
    _style_header(ws, 1, len(headers))
    for i, (name, info) in enumerate(projects.items(), 1):
        br = info["branches"]
        row = [i, name, "Yes", "", *br[:3]]
        row += [""] * (8 - len(row))
        ws.append(row)
        ws.cell(i + 1, 4).fill = _YELLOW      # Impacted Product = verify manually
    _autofit(ws, {"A": 6, "B": 34, "C": 14, "D": 18, "E": 24, "F": 24, "G": 24, "H": 20})
    ws.freeze_panes = "A2"


def _vuln_details_sheet(wb, current_q, prior_q, rows, history):
    ws = wb.create_sheet(_sheet_name(current_q, "VulnDetails"))

    # Row 1 banners: A-D blank, E:J = prior quarter, K:P = current quarter
    ws.append([""] * 16)
    for col, label in ((5, prior_q), (11, current_q)):
        ws.cell(1, col).value = label
        ws.merge_cells(start_row=1, start_column=col, end_row=1, end_column=col + 5)
        ws.cell(1, col).fill = _HEADER_FILL
        ws.cell(1, col).font = _BOLD
        ws.cell(1, col).alignment = _CENTER

    headers = ["No.", "Product", "Applicable for SQ Data", "Impacted Product",
               "Active Branch", "Critical", "High", "Medium", "Low", "Total (C+H)",
               "Active Branch", "Critical", "High", "Medium", "Low", "Total (C+H)"]
    ws.append(headers)
    _style_header(ws, 2, len(headers))

    r = 3
    first = r
    for idx, row in enumerate(rows, 1):
        prior = history.prior_for(prior_q, row["project"]) if history else None
        ws.cell(r, 1).value = idx
        ws.cell(r, 2).value = row["project"]
        ws.cell(r, 3).value = "Yes"
        ws.cell(r, 4).value = ""
        # prior block (E:J)
        if prior:
            ws.cell(r, 5).value = prior.get("Branch", "")
            ws.cell(r, 6).value = int(prior.get("Critical", 0) or 0)
            ws.cell(r, 7).value = int(prior.get("High", 0) or 0)
            ws.cell(r, 8).value = int(prior.get("Medium", 0) or 0)
            ws.cell(r, 9).value = int(prior.get("Low", 0) or 0)
        ws.cell(r, 10).value = f"=SUM(F{r}+G{r})"
        # current block (K:P)
        ws.cell(r, 11).value = row["branch"]
        ws.cell(r, 12).value = row["c"]
        ws.cell(r, 13).value = row["h"]
        ws.cell(r, 14).value = row["m"]
        ws.cell(r, 15).value = row["l"]
        ws.cell(r, 16).value = f"=SUM(L{r}+M{r})"
        for c in range(1, 17):
            ws.cell(r, c).border = _BORDER
        if row["c"]:
            ws.cell(r, 12).fill = _CRIT
        if row["h"]:
            ws.cell(r, 13).fill = _CRIT
        r += 1

    last = r - 1
    ws.cell(r, 1).value = "Total"
    ws.merge_cells(start_row=r, start_column=1, end_row=r, end_column=5)
    for c in list(range(6, 11)) + list(range(12, 17)):
        col = get_column_letter(c)
        ws.cell(r, c).value = f"=SUM({col}{first}:{col}{last})" if last >= first else 0
    ws.merge_cells(start_row=r, start_column=11, end_row=r, end_column=11)
    for c in range(1, 17):
        ws.cell(r, c).fill = _TOTAL_FILL
        ws.cell(r, c).font = _BOLD
        ws.cell(r, c).border = _BORDER

    _autofit(ws, {"A": 6, "B": 34, "C": 15, "D": 16,
                  "E": 20, "F": 9, "G": 7, "H": 8, "I": 7, "J": 11,
                  "K": 20, "L": 9, "M": 7, "N": 8, "O": 7, "P": 11})
    ws.freeze_panes = "A3"


def _ordered_rows(rows, project_order=None):
    """Sort rows for output.

    Default: alphabetical by project then branch (original behaviour).
    If `project_order` (a list of project names) is given, projects follow that
    order (case-insensitive, whitespace-trimmed). Any project present in the data
    but NOT in the list is appended afterwards alphabetically, so nothing is ever
    dropped. Branches stay sorted within each project.
    """
    if not project_order:
        return sorted(rows, key=lambda x: (x["project"].lower(), x["branch"].lower()))
    rank = {name.strip().lower(): i for i, name in enumerate(project_order) if name.strip()}
    tail = len(rank)                       # projects not in the list sort after all listed ones
    return sorted(rows, key=lambda x: (rank.get(x["project"].strip().lower(), tail),
                                       x["project"].lower(), x["branch"].lower()))


def build_workbook(studio: str, current_q: str, prior_q: str,
                   rows: list[dict], history, out_dir: str,
                   project_order: list[str] | None = None, tag: str = "") -> str:
    os.makedirs(out_dir, exist_ok=True)
    rows = _ordered_rows(rows, project_order)
    projects = _projects_in_order(rows)

    wb = Workbook()
    wb.remove(wb.active)
    _master_sheet(wb, current_q, projects)
    _active_branches_sheet(wb, current_q, projects)
    _vuln_details_sheet(wb, current_q, prior_q, rows, history)

    safe_q = current_q.replace("'", "").replace("/", "-")
    # Timestamp (down to the second) so every run writes a NEW file instead of
    # overwriting the last one -- avoids the "file is open in Excel" save failure
    # and lets you tell reports apart. Format YYYYMMDD_HHMMSS: filesystem-safe,
    # sorts chronologically in the folder.
    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    tag_part = f"_{tag}" if tag else ""
    path = os.path.join(out_dir, f"{studio}_{safe_q}{tag_part}_{timestamp}.xlsx")
    wb.save(path)
    return path


# --------------------------------------------------------------------------
# Read an existing report back so it can be re-sorted into a new project order
# --------------------------------------------------------------------------
class _PriorShim:
    """Minimal stand-in for scope.History, exposing just prior_for()."""

    def __init__(self, prior_map: dict):
        self._m = prior_map

    def prior_for(self, prior_q, project):
        return self._m.get(project)


def _find_sheet(wb, suffix):
    for ws in wb.worksheets:
        if ws.title.endswith(suffix):
            return ws
    return None


def _as_int(v):
    try:
        return int(v)
    except (TypeError, ValueError):
        return 0


def read_report(path: str):
    """Read a previously-generated report .xlsx back into
    (rows, prior_map, current_q, prior_q) so it can be regenerated in a new order.

    rows       : list of {project, key, branch, c, h, m, l} (current quarter)
    prior_map  : project -> {Branch, Critical, High, Medium, Low} (prior quarter)

    Raises ValueError if the file doesn't look like one of our reports.
    """
    wb = load_workbook(path, data_only=True)
    vd = _find_sheet(wb, "VulnDetails")
    if vd is None:
        raise ValueError("Not a SonarVulnReporter report (no VulnDetails sheet found).")
    master = _find_sheet(wb, "Master Sheet")

    prior_q = vd.cell(1, 5).value or ""
    current_q = vd.cell(1, 11).value or ""

    # project -> SonarQube key, from the Master sheet (col B = name, col F = key)
    keymap = {}
    if master is not None:
        r = 2
        while master.cell(r, 2).value not in (None, ""):
            keymap[str(master.cell(r, 2).value)] = str(master.cell(r, 6).value or "")
            r += 1

    rows, prior_map = [], {}
    r = 3
    while vd.cell(r, 1).value not in (None, "", "Total"):
        project = str(vd.cell(r, 2).value or "")
        rows.append({
            "project": project,
            "key": keymap.get(project, ""),
            "branch": str(vd.cell(r, 11).value or ""),
            "c": _as_int(vd.cell(r, 12).value),
            "h": _as_int(vd.cell(r, 13).value),
            "m": _as_int(vd.cell(r, 14).value),
            "l": _as_int(vd.cell(r, 15).value),
        })
        prior_map[project] = {
            "Branch": str(vd.cell(r, 5).value or ""),
            "Critical": _as_int(vd.cell(r, 6).value),
            "High": _as_int(vd.cell(r, 7).value),
            "Medium": _as_int(vd.cell(r, 8).value),
            "Low": _as_int(vd.cell(r, 9).value),
        }
        r += 1
    return rows, prior_map, str(current_q), str(prior_q)


def resort_report_file(studio: str, src_path: str, project_order: list[str],
                       out_dir: str) -> str:
    """Read an existing report, re-order its projects per `project_order`, and
    write a NEW workbook (tagged '_sorted'). Returns the new file path."""
    rows, prior_map, current_q, prior_q = read_report(src_path)
    return build_workbook(studio, current_q, prior_q, rows,
                          _PriorShim(prior_map), out_dir,
                          project_order=project_order, tag="sorted")
